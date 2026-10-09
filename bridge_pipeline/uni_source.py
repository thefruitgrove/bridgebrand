from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path

from .targets import load_targets


SOURCE_URL = "https://www.data.go.kr/data/15118998/fileData.do"
OUTPUT_FIELDS = [
    "institution", "external_id", "school_type", "region_group", "competition_rate",
    "freshman_fill_rate", "international_student_rate", "program_distinctiveness",
    "specialization_concentration", "retention_rate", "dropout_rate",
    "scholarship_per_student", "education_spend_per_student", "dorm_capacity_rate",
    "faculty_capacity_rate", "library_resources_per_student", "employment_rate",
    "research_per_faculty", "industry_revenue_per_student", "tech_transfer_income",
    "disclosure_completeness", "financial_stability", "education_condition_trend",
    "freshness_days", "evidence_url", "note",
]


def compact(value: str | None) -> str:
    return re.sub(r"\s+", "", str(value or ""))


def numeric(value):
    if value is None:
        return None
    text = str(value).replace(",", "").replace("%", "").strip()
    if not text or text in {"-", "N/A"}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _find(row: dict, prefix: str):
    key = next((key for key in row if compact(key).startswith(compact(prefix))), None)
    return row.get(key) if key else None


def convert(rows: list[dict]) -> tuple[list[dict], list[str]]:
    # Prefer the main campus when a university has multiple disclosure rows.
    ordered = sorted(rows, key=lambda row: 0 if row.get("본분교명") == "본교" else 1)
    by_name = {}
    for row in ordered:
        by_name.setdefault(compact(row.get("학교명")), row)

    output: list[dict] = []
    missing: list[str] = []
    for target in (row for row in load_targets() if row["domain"] == "UNI"):
        candidates = [target["name"], *(target.get("aliases") or [])]
        source = next((by_name[compact(name)] for name in candidates if compact(name) in by_name), None)
        if source is None:
            missing.append(target["name"])
            continue

        students = numeric(_find(source, "재학생"))
        international = numeric(_find(source, "외국인 학생 수"))
        education_cost = numeric(_find(source, "학생 1인당 교육비"))
        tuition = numeric(_find(source, "평균 등록금"))
        raw_indicators = [
            _find(source, "신입생 경쟁률"), _find(source, "신입생 충원율"),
            _find(source, "취업률"), _find(source, "외국인 학생 수"),
            _find(source, "전임교원 1인당 학생 수"), _find(source, "전임교원 확보율(학생정원"),
            _find(source, "전임 교원 확보율(재학생"), _find(source, "전임교원 강의 담당 비율"),
            _find(source, "학생 1인당 연간 장학금"), education_cost,
            _find(source, "기숙사 수용률"), _find(source, "학생 1인당 도서 자료 수"),
        ]
        present = sum(numeric(value) is not None for value in raw_indicators)
        region = str(source.get("지역명") or "")
        output.append({
            "institution": source["학교명"],
            "external_id": target["external_id"],
            "school_type": "4년제" if target["subcategory"] == "U60" else "전문대",
            "region_group": "수도권" if region in {"서울", "경기", "인천"} else "비수도권",
            "competition_rate": _find(source, "신입생 경쟁률"),
            "freshman_fill_rate": _find(source, "신입생 충원율"),
            "international_student_rate": round(100 * international / students, 4) if students and international is not None else "",
            "scholarship_per_student": _find(source, "학생 1인당 연간 장학금"),
            "education_spend_per_student": education_cost if education_cost is not None else "",
            "dorm_capacity_rate": _find(source, "기숙사 수용률"),
            "faculty_capacity_rate": _find(source, "전임 교원 확보율(재학생"),
            "library_resources_per_student": _find(source, "학생 1인당 도서 자료 수"),
            "employment_rate": _find(source, "취업률"),
            "disclosure_completeness": round(100 * present / len(raw_indicators), 2),
            "financial_stability": round(100 * education_cost / tuition, 4) if education_cost is not None and tuition else "",
            "freshness_days": "0",
            "evidence_url": SOURCE_URL,
            "note": "대학알리미 최신 대학주요정보(지표별 2025~2026 공시)",
        })
    return output, missing


def main(argv=None):
    parser = argparse.ArgumentParser(description="Convert AcademyInfo university overview XLSX to BRIDGE UNI input")
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args(argv)
    from openpyxl import load_workbook

    worksheet = load_workbook(args.input, read_only=True, data_only=True).active
    values = worksheet.iter_rows(values_only=True)
    headers = [str(value or "") for value in next(values)]
    rows = [dict(zip(headers, row)) for row in values]
    converted, missing = convert(rows)
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=OUTPUT_FIELDS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(converted)
    print(f"matched {len(converted)}/100 UNI targets; missing={','.join(missing) or '-'}")


if __name__ == "__main__":
    main()
